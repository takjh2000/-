"""기존 엑셀(보드게임 대여 시스템) 재고 데이터를 DB로 1회성 이관하는 스크립트.

사용법:
    python -m app.seed.import_excel "/Users/takjh2000/Downloads/보드게임 대여 시스템-2.xlsx"

대여기록/연체기록은 옮기지 않는다. 기존 대여자가 새 시스템의 회원 계정과
연결되어 있지 않아 그대로 옮기면 무결성이 깨지기 때문에, 재고(게임 목록)만
이관하고 대여 이력은 새 시스템에서부터 새로 쌓는다.
"""
import sys

import openpyxl

from ..database import SessionLocal, Base, engine
from ..models import Game, GameCategory

SHEET_CATEGORY_MAP = {
    "보드게임_현황": GameCategory.boardgame,
    "크라임씬_현황": GameCategory.crimescene,
}


def _clean(value):
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def import_games(xlsx_path: str) -> None:
    Base.metadata.create_all(bind=engine)
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    db = SessionLocal()
    created, skipped = 0, 0
    try:
        for sheet_name, category in SHEET_CATEGORY_MAP.items():
            if sheet_name not in wb.sheetnames:
                continue
            ws = wb[sheet_name]
            for row in ws.iter_rows(min_row=2, values_only=True):
                name = _clean(row[0])
                if not name:
                    continue
                owner = _clean(row[1])
                total_quantity = int(row[2]) if row[2] else 1
                notes = _clean(row[5]) if len(row) > 5 else None

                exists = (
                    db.query(Game)
                    .filter(Game.name == name, Game.category == category)
                    .first()
                )
                if exists:
                    skipped += 1
                    continue

                db.add(
                    Game(
                        name=name,
                        category=category,
                        owner=owner,
                        total_quantity=total_quantity,
                        notes=notes,
                    )
                )
                created += 1
        db.commit()
    finally:
        db.close()

    print(f"완료: {created}건 생성, {skipped}건 중복 건너뜀")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        print("사용법: python -m app.seed.import_excel <xlsx 경로>")
        sys.exit(1)
    import_games(sys.argv[1])
